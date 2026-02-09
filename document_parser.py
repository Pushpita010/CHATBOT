import os
from PyPDF2 import PdfReader
from docx import Document
import openpyxl


def extract_text_from_file(file_obj, filename):
    """
    Extract text from PDF, Word (.docx), or Excel (.xlsx) files.

Args:
            file_obj: File-like object (BytesIO) containing file data
            filename: Original filename to determine file type
    """
    ext = os.path.splitext(filename)[1].lower()
    if ext == '.pdf':
        return extract_text_from_pdf(file_obj)
    elif ext == '.docx':
        return extract_text_from_docx(file_obj)
    elif ext == '.xlsx':
        return extract_text_from_xlsx(file_obj)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def extract_text_from_pdf(file_obj):
    text = []
    reader = PdfReader(file_obj)
    for page in reader.pages:
        text.append(page.extract_text() or '')
    return '\n'.join(text)


def extract_text_from_docx(file_obj):
    doc = Document(file_obj)
    text = [para.text for para in doc.paragraphs]
    return '\n'.join(text)


def extract_text_from_xlsx(file_obj):
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    text = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = [str(cell) if cell is not None else '' for cell in row]
            text.append('\t'.join(row_text))
    return '\n'.join(text)
